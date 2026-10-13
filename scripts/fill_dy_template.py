#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把「非抖音」素材包（淘宝/拼多多 sheet）的商品数据，灌进**抖店官方模板**的「抖音」sheet。

为什么需要它：抖店各类目的官方素材包模板（商品信息.xlsx）随类目变化，
而用户的素材往往只有「拼多多」「淘宝」sheet 有数据。本脚本以官方模板为
**列定义唯一来源**（R2 字段名 / R3 SKU 子表头 / 选项_dy 允许值），
只把数据搬进去，不改模板结构。

用法：
  python fill_dy_template.py \
    --template "<官方模板>\\商品信息.xlsx" \
    --src-xlsx "<素材包>\\发品excel\\ZDT-002.xlsx" \
    --product-id ZDT-002 \
    --out "<成品包>\\商品信息.xlsx" \
    --attr "材质=聚酯纤维（涤纶）" --attr "风格=现代简约"

产出：<out> 一份填好的官方模板（保留原 sheet 与下拉），随后可直接跑
      read_excel.py / build_product.py。
"""
import argparse
import json
import os
import re
import shutil
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

import openpyxl

# 淘宝/拼多多 字段名 → 抖音 字段名（只做「同名或语义等价」的搬运，不臆造值）
ATTR_ALIAS = {
    "表面材质": "材质",
    "主体材质": "材质",
    "底背材质": "材质",
    "底部材质": "材质",
    "饰面材质": "材质",
    "编织工艺": "工艺",
    "工艺": "工艺",
    "主图案类型": "图案",
    "图案": "图案",
    "风格": "风格",
    "形状": "形状",
    "shape": "形状",
    "产地": "产地",
    "origin": "产地",
    "货号": "货号",
    "售卖方式": "售卖方式",
    "适用空间": "适用空间",
    "适用场景": "适用场景",
    "清洗类型": "清洗类型",
    "清洁方式": "清洁方式",
    "功能特点": "功能",
    "功能": "功能",
    "厚度": "厚度",
    "品牌": "品牌",
}

# 抖音 SKU 子表头 → 我们希望从源 sheet 取的列（源列按 R3 子表头找）
SKU_SOURCE = {
    "价格": ["价格", "拼单价", "单买价", "一口价", "skuPrice"],
    "现货库存": ["库存", "现货库存", "数量", "skuStock"],
    "商家编码": ["规格编码", "商家编码", "skuCode", "outerId"],
    "条形码": ["条形码", "商品条形码", "skuBarcode"],
    "规格图片文件名": ["预览图文件名", "规格图片文件名", "规格图片"],
}


def norm(s):
    if s is None:
        return None
    t = re.sub(r"[\s★☆]+", "", str(s))
    return t or None


def col_labels(ws, ncols):
    """返回 [{col, label, label2, label3}]：label 优先取 R3 子表头（与 read_excel 一致）。"""
    out = []
    for c in range(1, ncols + 1):
        l2, l3 = ws.cell(2, c).value, ws.cell(3, c).value
        lab = None
        for cand in (l3, l2):
            if cand is not None and str(cand).strip():
                lab = re.sub(r"★", "", str(cand)).strip()
                break
        out.append({"col": c, "label": lab, "label2": norm(l2), "label3": norm(l3)})
    return out


def find_col(labels, want, prefer_sku=False):
    w = norm(want)
    hits = []
    for L in labels:
        if L["label3"] == w:
            hits.append((2, L))
        elif L["label2"] == w:
            hits.append((1, L))
    if not hits:
        return None
    hits.sort(key=lambda x: -x[0])
    return hits[0][1]["col"]


def first_data_row(ws, ncols):
    """源 sheet 的第一个非空数据行（淘宝/拼多多都是 R5 起）"""
    for r in range(5, ws.max_row + 1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, ncols + 1)):
            return r
    return None


def strip_brackets(s):
    """去掉括号及其内容（圆括号/方括号、全半角都算）。"""
    return re.sub(r"[（(【\[][^）)】\]]*[）)】\]]", "", str(s)).strip()


def brackets_to_dash(s):
    """把括号**转成连字符**（保留括号里的信息），并收敛重复的 `-`。

    例：`白色#【30倍殿堂品质】9.0纳米旗舰暴粘款` → `白色#-30倍殿堂品质-9.0纳米旗舰暴粘款`
    """
    t = re.sub(r"\s*[（(【\[]\s*([^）)】\]]*?)\s*[）)】\]]\s*", r"-\1-", str(s))
    return re.sub(r"-{2,}", "-", t).strip("-").strip()


def transform_spec_value(value, mode):
    """按 mode 处理规格值里的括号：'strip' 剥掉 / 'dash' 转连字符。

    ⚠️ 判据是「**剥完会不会撞**」，不是「括号里有没有字」：
      · 会撞（括号是唯一区分信息）→ 必须 'dash'
        PE3 实测：`白色（2.0版/标准遮丑）`/`白色（3.0版/强胶暴粘不透底）` 剥完都成 `白色`，12 值塌成 4。
      · 不会撞（括号后面还有区分文字 / 前缀在所有值里都一样）→ 两者都行
        GPE-002 实测：`白色#【10倍非凡品质】5.0标准款` 剥完是 `白色#5.0标准款`，**仍是 3 个不同值**；
        GPVC-001 实测：`【真格栅】…` 8 个值前缀全同，剥完互相区分 → 取 strip 让选项文字更短。
    （我最初把「取 dash」的理由写成「不转就会撞」，被 scripts/test_dy_logic.py 抓出来了。）
    """
    if not value:
        return value
    out = brackets_to_dash(value) if mode == "dash" else strip_brackets(value)
    return out or value


def merge_combo_code(dim1, dim2, prefix):
    """多色拼接（规格类型1 的值形如「棕色+卡其-10片」）的编码模板：

        {前缀}-{尺码}-{色A}+{色B}各{N/2}片          N = 该值里的**总**片数

    例：`棕色+卡其-10片` @ `40*40cm` → `自粘地毯-40*40-棕色+卡其各5片`

    **不要**从源表的「规格编码」推导：拼接行是两条单色码用 `+` 硬拼的，
    而且片数自相矛盾（6片 行是 `5片+5片`、10片 行是 `10片+10片`）。
    拼不出「各N片」（总片数为奇数）时返回 None，保持源编码不动。
    """
    v = str(dim1 or "")
    if "+" not in v:
        return None
    combo, _, pack = v.rpartition("-")
    if not combo or not pack:
        return None
    m = re.search(r"(\d+)", pack)
    if not m:
        return None
    n = int(m.group(1))
    if n <= 0 or n % 2:
        return None
    size = re.sub(r"(厘米|cm|CM)$", "", str(dim2 or "")).strip()
    if not size:
        return None
    return "%s-%s-%s各%d片" % (prefix, size, combo, n // 2)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True, help="抖店官方模板 商品信息.xlsx（只读列定义）")
    ap.add_argument("--src-xlsx", required=True, help="数据来源 xlsx（含 淘宝/拼多多 sheet）")
    ap.add_argument("--out", required=True, help="填好后的输出 xlsx")
    ap.add_argument("--product-id", required=True)
    ap.add_argument("--sku-sheet", default="拼多多", help="SKU/价格来源 sheet")
    ap.add_argument("--attr-sheet", default="淘宝", help="产品级属性来源 sheet（可留空字符串跳过）")
    ap.add_argument("--price-source", required=True,
                    help="价格取源哪一列（拼单价/单买价/一口价）。**必填且没有默认值** —— "
                         "抖店是「售卖价」（无拼单概念），而源素材包常同时给两列（实测 ZDT-002: "
                         "拼单价 13.6~82 vs 单买价 14.6~83）。口径必须由用户确认后显式传入（门禁 42）")
    ap.add_argument("--brand", default=None,
                    help="品牌。抖店品牌必填且必须是你店已申请/已授权的，源数据常与实际不一致"
                         "（实测：源写「源品牌A」「源品牌B」，实际用「示例品牌」）→ 问用户后在这里传（门禁 43）。"
                         "不传则用源数据里的品牌，报告里会标明来源")
    ap.add_argument("--dim-order", default="颜色,尺码",
                    help="抖音 规格类型1/2 依次取源 sheet 的哪些维度（逗号分隔）；"
                         "顺序要紧：规格图只能挂第一组，出图维度必须排最前")
    ap.add_argument("--strip-spec-parens", action="store_true",
                    help="去掉规格值里的括号补充说明（圆括号/方括号都算，含全角）。抖店智能识别会把 "
                         "`（…）`/`【…】` 当补充说明剥掉，源值写成「3米装（卧室书房首选）」时 3 个值会塌成"
                         "「3米装」1 个（PE3 实测 12 值塌成 4；MZB-015 的「300厘米*60厘米【不伤爪】」同理）")
    ap.add_argument("--bracket-to-dash", action="store_true",
                    help="把括号**转成连字符**而不是丢掉：`白色#【30倍殿堂品质】9.0纳米旗舰暴粘款` → "
                         "`白色#-30倍殿堂品质-9.0纳米旗舰暴粘款`。**剥完会撞时必须用这个** —— "
                         "PE3 实测「白色（2.0版）/白色（3.0版）」剥完都成「白色」，12 个值塌成 4 个；"
                         "剥完不撞但括号里是有效标签时也建议用（GPE-002 的「10倍/20倍/30倍品质」）。"
                         "两种模式都可行时，跑 scripts/test_dy_logic.py 或看 dim1_unique/dim2_unique 是否等于预期值数。"
                         "两者同时给时本项优先")
    ap.add_argument("--attr", action="append", default=[], help="覆盖/追加属性：名称=值（可重复）")
    ap.add_argument("--title", default=None)
    ap.add_argument("--stock", type=int, default=None, help="覆盖库存（默认取源）")
    ap.add_argument("--combo-code-merge", action="store_true",
                    help="多色拼接（规格类型1 的值含 `+`）的「商家编码」改写成 "
                         "「{前缀}-{尺码}-{色A}+{色B}各{总片数/2}片」。"
                         "例：棕色+卡其-10片 @ 40*40cm → 自粘地毯-40*40-棕色+卡其各5片。"
                         "源表里拼接行的编码是两条单色码用 + 硬拼的、且片数自相矛盾，不能直接用")
    ap.add_argument("--combo-code-prefix", default=None,
                    help="拼接编码的前缀；默认识别源编码里 `-` 之前的那一段")
    a = ap.parse_args()

    tpl = openpyxl.load_workbook(a.template)          # 保留公式/下拉
    if "抖音" not in tpl.sheetnames:
        raise SystemExit("模板里没有「抖音」sheet")
    dws = tpl["抖音"]
    dlabels = col_labels(dws, dws.max_column)

    src = openpyxl.load_workbook(a.src_xlsx, data_only=True)
    sname = a.sku_sheet if a.sku_sheet in src.sheetnames else src.sheetnames[2]
    sws = src[sname]
    slabels = col_labels(sws, sws.max_column)
    r0 = first_data_row(sws, sws.max_column)
    if r0 is None:
        raise SystemExit("源 sheet「%s」没有数据行" % sname)

    # ---- 源列定位 ----
    def src_col(cands, sheet_labels, want_sku=False):
        for w in cands:
            c = find_col(sheet_labels, w)
            if c:
                return c
        return None

    dim_src = {}
    for d in [x.strip() for x in a.dim_order.split(",") if x.strip()]:
        c = find_col(slabels, d)
        if not c:
            raise SystemExit("源 sheet 找不到维度「%s」，可选：%s"
                             % (d, [L["label"] for L in slabels if L["label"]]))
        dim_src[d] = c
    price_c = src_col([a.price_source, "价格", "一口价"], slabels)
    stock_c = src_col(["库存", "现货库存", "数量"], slabels)
    code_c = src_col(["规格编码", "商家编码", "商品编码"], slabels)
    img_c = src_col(SKU_SOURCE["规格图片文件名"], slabels)
    title_c = src_col(["商品标题", "宝贝标题", "title"], slabels)
    if not price_c or not stock_c:
        raise SystemExit("源 sheet 缺价格或库存列")

    # ---- SKU 行（按源数据顺序；同一维度组合只保留一次）----
    dim_names = [x.strip() for x in a.dim_order.split(",") if x.strip()]
    stripped = []
    skus = []
    seen = set()
    for r in range(r0, sws.max_row + 1):
        vals = {}
        for d, c in dim_src.items():
            v = sws.cell(r, c).value
            vals[d] = str(v).strip() if v not in (None, "") else None
            if a.strip_spec_parens or a.bracket_to_dash:
                before = vals[d]
                vals[d] = transform_spec_value(vals[d], "dash" if a.bracket_to_dash else "strip")
                if vals[d] != before:
                    stripped.append("%s → %s" % (before, vals[d]))
        if not any(vals.values()):
            continue
        key = tuple(vals[d] for d in dim_names)
        if key in seen:
            continue
        seen.add(key)
        skus.append({
            "dims": vals,
            "price": sws.cell(r, price_c).value,
            "stock": a.stock if a.stock is not None else sws.cell(r, stock_c).value,
            "code": (str(sws.cell(r, code_c).value).strip()
                     if code_c and sws.cell(r, code_c).value not in (None, "") else None),
            "img": (str(sws.cell(r, img_c).value).strip()
                    if img_c and sws.cell(r, img_c).value not in (None, "") else None),
        })

    # ---- 多色拼接的「商家编码」合并 ----
    # 源表里拼接行写的是**两条单色码用 + 硬拼**，且片数自相矛盾
    # （「棕色+卡其-6片」行的编码是 `…-5片+…-5片`，「-10片」行是 `…-10片+…-10片`），
    # 所以只认「规格类型1 的值」里的总片数，改成「…各{总片数/2}片」。
    combo_merge = []
    if a.combo_code_merge:
        prefixes = {str(s["code"]).split("-")[0] for s in skus if s.get("code")}
        prefix = a.combo_code_prefix or (prefixes.pop() if len(prefixes) == 1 else None)
        if not prefix:
            print("⚠ --combo-code-merge：推不出编码前缀（源编码前缀不唯一），跳过合并")
        else:
            for s in skus:
                merged = merge_combo_code(s["dims"].get(dim_names[0]),
                                          s["dims"].get(dim_names[1]) if len(dim_names) > 1 else None,
                                          prefix)
                if merged:
                    if merged != s["code"]:
                        combo_merge.append({"color": s["dims"].get(dim_names[0]),
                                            "size": s["dims"].get(dim_names[1]) if len(dim_names) > 1 else None,
                                            "from": s["code"], "to": merged})
                    s["code"] = merged

    # ---- 产品级属性 ----
    attrs = {}
    mapping_note = []
    if a.attr_sheet and a.attr_sheet in src.sheetnames:
        aws = src[a.attr_sheet]
        alabels = col_labels(aws, aws.max_column)
        ar0 = first_data_row(aws, aws.max_column)
        for L in alabels:
            if not L["label"]:
                continue
            target = ATTR_ALIAS.get(L["label"])
            if not target or target in attrs:
                continue
            # 目标字段必须在抖音模板里存在
            if not find_col(dlabels, target):
                continue
            v = aws.cell(ar0, L["col"]).value
            if v in (None, ""):
                continue
            attrs[target] = str(v).strip()
            mapping_note.append("%s.%s → %s = %s" % (a.attr_sheet, L["label"], target, attrs[target]))

    for kv in a.attr:
        if "=" not in kv:
            raise SystemExit("--attr 需要 名称=值：%s" % kv)
        k, v = kv.split("=", 1)
        attrs[k.strip()] = v.strip()
        mapping_note.append("override → %s = %s" % (k.strip(), v.strip()))

    # 品牌（门禁 43）：显式传入即覆盖源数据，并在报告里标明「源数据是什么」——
    # 让「源写源品牌A、我们填示例品牌」这种决定留下痕迹，而不是悄悄换掉。
    brand_source = attrs.get("品牌")
    if a.brand:
        attrs["品牌"] = a.brand
        mapping_note.append("override → 品牌 = %s（源数据是 %s）" % (a.brand, brand_source or "（空）"))
    elif brand_source:
        mapping_note.append("品牌 = %s（直接取自源数据，未显式覆盖 —— 确认过是你店已申请的品牌吗？）" % brand_source)
    else:
        print("⚠ 源数据没有品牌，且没传 --brand —— 抖店品牌是必填，请先问用户（门禁 43）")

    title = a.title or (str(sws.cell(r0, title_c).value).strip() if title_c and sws.cell(r0, title_c).value else None)
    if not title:
        raise SystemExit("没有标题（用 --title 指定）")

    # ---- 写入抖音 sheet ----
    def w(col, row, val):
        if col and val not in (None, ""):
            dws.cell(row, col).value = val

    # 产品标识列名各模板不同：常见「商品标识」，也有模板叫「商品名称」（猫抓板类目实测）
    c_pid = find_col(dlabels, "商品标识") or find_col(dlabels, "商品名称") or find_col(dlabels, "__productId")
    c_title = find_col(dlabels, "商品标题")
    c_dim1 = find_col(dlabels, "规格类型1")
    c_dim2 = find_col(dlabels, "规格类型2")
    c_dim3 = find_col(dlabels, "规格类型3")
    c_img = find_col(dlabels, "规格图片文件名")
    c_price = find_col(dlabels, "价格")
    c_stock = find_col(dlabels, "现货库存")
    c_code = find_col(dlabels, "商家编码")
    c_barcode = find_col(dlabels, "条形码")
    for need, c in [("规格类型1", c_dim1), ("价格", c_price), ("现货库存", c_stock)]:
        if not c:
            raise SystemExit("模板里找不到列「%s」" % need)

    w(c_pid, 5, a.product_id)
    w(c_title, 5, title)
    for k, v in attrs.items():
        c = find_col(dlabels, k)
        if c:
            w(c, 5, v)
        else:
            print("⚠ 抖音模板没有「%s」字段，跳过（值=%s）" % (k, v))

    for i, s in enumerate(skus):
        r = 5 + i
        w(c_dim1, r, s["dims"].get(dim_names[0]))
        if len(dim_names) > 1 and c_dim2:
            w(c_dim2, r, s["dims"].get(dim_names[1]))
        if len(dim_names) > 2 and c_dim3:
            w(c_dim3, r, s["dims"].get(dim_names[2]))
        w(c_img, r, s["img"])
        w(c_price, r, s["price"])
        w(c_stock, r, s["stock"])
        w(c_code, r, s["code"])

    # ★ 残留守卫：这份模板是从**上一次链接**导出的，R5 里可能还留着旧值
    #   （实测 猫抓板模板：材质=「聚脂钎维」、商品详情=旧富文本、参考价设置=100）。
    #   我们只写「商品标识/标题/属性/SKU 列」，其它列原样不动 —— 但 build_product 会把其中一部分
    #   当数据读（实测「参考价设置=100」被读成 ref_price=100）；而 ★必填属性若因别名缺失**没写进去**，
    #   残留值就会冒充成品数据（实测 材质 差点被旧值顶替）。所以这里把「我们没写却非空」的列报出来。
    written = {c for c in [c_pid, c_title, c_dim1, c_dim2, c_dim3, c_img, c_price, c_stock,
                           c_code, c_barcode] if c}
    for k in attrs:
        written.add(find_col(dlabels, k))
    lab_of = {L["col"]: L["label"] for L in dlabels}
    resid = []
    for c in range(1, dws.max_column + 1):
        if c in written:
            continue
        v = dws.cell(5, c).value
        if v in (None, ""):
            continue
        resid.append("%s=%s" % (lab_of.get(c) or ("列%d" % c), str(v)[:36]))
    if resid:
        print("⚠ 模板 R5 有 %d 处**我们没写**的列已有值（模板预填 / 旧链接残留）：%s"
              % (len(resid), "; ".join(resid[:8])))
        print("   多数无害（流程会在页面上重设这些字段，主图/详情图另有目录来源），"
              "但 ★必填属性 与 build_product 会读的字段要人工核对。")

    os.makedirs(os.path.dirname(os.path.abspath(a.out)) or ".", exist_ok=True)
    tpl.save(a.out)

    n1 = len({s["dims"][dim_names[0]] for s in skus})
    report = {
        "out": a.out,
        "product_id": a.product_id,
        "title": title,
        "title_chars": len(title),
        "title_hanzi": len(re.findall(r"[\u4e00-\u9fff]", title)),
        "sku_sheet": sname,
        "sku_rows": len(skus),
        "dim_order": dim_names,
        "dim1_unique": n1,
        "dim2_unique": len({s["dims"][dim_names[1]] for s in skus}) if len(dim_names) > 1 else 0,
        "price_source": a.price_source,
        "price_min": min(s["price"] for s in skus),
        "price_max": max(s["price"] for s in skus),
        "images_referenced": len({s["img"] for s in skus if s["img"]}),
        "combo_code_merged": len(combo_merge),
        "combo_code_samples": combo_merge[:3],
        "spec_values_paren_stripped": sorted(set(stripped))[:8],
        "spec_values_paren_stripped_n": len(set(stripped)),
        "code_unique": len({str(s["code"]) for s in skus if s["code"]}),
        "attributes": attrs,
        "mapping": mapping_note,
    }
    print(json.dumps(report, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
