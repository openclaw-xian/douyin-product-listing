#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
解析抖店官方素材包 Excel（商品信息.xlsx）的「抖音」sheet。

结构（2026-09-30 实测）：
  R1  平台标题：抖音 | platform=dy | 家装建材 > 墙面材料 > 墙纸 | cateId=1000004860
  R2  字段名（66 列）
  R3  SKU 子表头：规格类型1 / 规格类型2 / 规格类型3 / 规格图片文件名 / ★价格 / ★现货库存 / 商家编码 / 条形码
  R4  JSON 定义：{"key":"...","required":true,"assetPath":["主图"],"kind":"skuSpec","skuRole":"specImage"}
  R5+ 数据行（每个 SKU 一行；产品级字段只在第 5 行写一次）

用法：
  python read_excel.py "C:/path/商品信息.xlsx"
  exit 0=ok  1=路径/解析硬错误  2=预检未过
"""
import json
import os
import re
import sys

# Windows 下 stdout 默认是 cp936(GBK)：中文会被转成 ? / U+FFFD，
# 调用方（build_product.py）再按 utf-8 解就拿到乱码。入口处强制 UTF-8。
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

try:
    import openpyxl
except ImportError:
    print(json.dumps({"ok": False, "error": "openpyxl_missing",
                      "hint": "用装了 openpyxl 的解释器（MIMO_PYTHON / config.python）"}, ensure_ascii=False))
    sys.exit(1)

PLATFORMS = ("抖音", "拼多多", "淘宝")
IGNORE_SHEETS = ("填写说明", "素材包设置")

DEFAULT_DIMS = ["规格类型1", "规格类型2", "规格类型3"]
MAX_SPEC_DIMS = 3


def pick_platform_sheet(wb):
    """抖音 sheet 优先。

    找不到「抖音」时**不再静默兜底**到别的平台 sheet：素材包可能只有
    「拼多多」sheet，拿它当抖店数据会一路错填到后台且毫无提示。
    返回 (sheet_name, error)。
    """
    if "抖音" in wb.sheetnames:
        return "抖音", None
    others = [n for n in wb.sheetnames
              if n not in IGNORE_SHEETS and not n.startswith("选项")]
    if others:
        return None, {"error": "no_douyin_sheet", "found_sheets": others,
                      "hint": "该素材包没有「抖音」sheet，不能当抖店数据用；"
                              "拼多多请用 pdd-product-listing"}
    return None, {"error": "no_douyin_sheet", "found_sheets": list(wb.sheetnames)}


def parse_title_row(v):
    """从 R1 解析 category / cateId。"""
    s = str(v or "")
    category = None
    cate_id = None
    m = re.search(r"cateId\s*=\s*(\d+)", s)
    if m:
        cate_id = m.group(1)
    parts = [p.strip() for p in s.split("|")]
    for p in parts:
        if ">" in p:
            category = p
            break
    return category, cate_id


def norm_label(s):
    """表头带 ★ 前缀（如「★ 商品标题」）→ 归一成「商品标题」"""
    if s is None:
        return None
    t = str(s).replace("★", "").strip()
    return t or None


def key_label(s):
    """归一成用于**比对**的键：去掉 ★ 与所有空白。

    官方模板的 R3 子表头可能写成「规格类型 1」（内部带空格），
    直接拿 norm_label 的结果去比 DEFAULT_DIMS 会匹配不上，
    维度识别随之落空、SKU 全部解析成 0 行。
    """
    t = norm_label(s)
    if t is None:
        return None
    return re.sub(r"\s+", "", t) or None


def load_options(wb):
    """选项_dy：属性名 → 允许值列表。"""
    sheet = None
    for name in wb.sheetnames:
        if name.startswith("选项"):
            sheet = name
            break
    if not sheet:
        return {}
    ws = wb[sheet]
    out = {}
    for c in range(1, ws.max_column + 1):
        head = ws.cell(1, c).value
        if head is None or not str(head).strip():
            continue
        vals = []
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip():
                vals.append(str(v).strip())
        out[str(head).strip()] = vals
    return out


def main():
    if len(sys.argv) < 2:
        print(json.dumps({"ok": False, "error": "usage: read_excel.py <xlsx>"}, ensure_ascii=False))
        sys.exit(1)
    xlsx = sys.argv[1]
    if not os.path.isfile(xlsx):
        print(json.dumps({"ok": False, "error": "file_not_found", "path": xlsx}, ensure_ascii=False))
        sys.exit(1)

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    sheet_name, sheet_err = pick_platform_sheet(wb)
    if sheet_err:
        print(json.dumps({"ok": False, **sheet_err}, ensure_ascii=False, indent=1))
        sys.exit(1)
    ws = wb[sheet_name]

    category, cate_id = parse_title_row(ws.cell(1, 1).value)

    # 列定义
    cols = []
    for c in range(1, ws.max_column + 1):
        label2 = ws.cell(2, c).value
        label3 = ws.cell(3, c).value
        raw4 = ws.cell(4, c).value
        meta = {}
        if raw4:
            try:
                meta = json.loads(raw4)
            except Exception:
                meta = {}
        label = None
        for cand in (label3, label2):
            if cand is not None and str(cand).strip():
                label = norm_label(cand)
                break
        if not label and not meta.get("key"):
            continue
        cols.append({
            "col": c,
            "label": label,
            "label2": norm_label(label2),
            "raw_label": str(label3 or label2 or "").strip() or None,
            "key": meta.get("key"),
            "required": bool(meta.get("required")),
            "kind": meta.get("kind"),
            "assetPath": meta.get("assetPath"),
            "skuRole": meta.get("skuRole"),
        })

    sku_cols = [c for c in cols if c.get("kind") == "skuSpec"]
    if not sku_cols:
        print(json.dumps({"ok": False, "error": "no_sku_columns", "sheet": sheet_name}, ensure_ascii=False))
        sys.exit(1)

    # 规格维度：SKU 子表头里非 fillColumn 的（规格类型1/2/3）
    # 用 key_label 比对：官方模板可能写成「规格类型 1」（内部空格）
    dim_cols = [c for c in sku_cols if c.get("skuRole") != "fillColumn"
                and key_label(c.get("label")) in DEFAULT_DIMS]
    dim_cols.sort(key=lambda x: DEFAULT_DIMS.index(key_label(x["label"])))
    if not dim_cols:
        fallback = [c for c in sku_cols if c.get("skuRole") != "fillColumn"][:MAX_SPEC_DIMS]
        # 兜底只在「看着像规格类型N」时才认，否则宁可报错也不乱取列
        dim_cols = [c for c in fallback
                    if re.fullmatch(r"规格类型\s*\d", key_label(c.get("label")) or "")]
        if not dim_cols:
            print(json.dumps({
                "ok": False, "error": "spec_dim_columns_not_found",
                "sku_labels": [c.get("label") for c in sku_cols],
                "hint": "SKU 子表头没认出「规格类型1/2/3」；确认素材包 R3 行格式",
            }, ensure_ascii=False, indent=1))
            sys.exit(1)

    def find_sku_col(label_text):
        for c in sku_cols:
            if c.get("label") and label_text in c["label"]:
                return c
        return None

    def find_sku_col_norm(label_text):
        """去空白后查找（表头可能带内部空格）"""
        want = re.sub(r"\s+", "", label_text)
        for c in sku_cols:
            got = key_label(c.get("label"))
            if got and want in got:
                return c
        return find_sku_col(label_text)

    col_price = find_sku_col_norm("价格")
    col_stock = find_sku_col_norm("库存")
    col_code = find_sku_col_norm("商家编码")
    col_barcode = find_sku_col_norm("条形码")
    if not col_price or not col_stock:
        print(json.dumps({
            "ok": False, "error": "price_or_stock_column_missing",
            "sku_labels": [c.get("label") for c in sku_cols],
        }, ensure_ascii=False, indent=1))
        sys.exit(1)
    col_spec_img = next((c for c in sku_cols if c.get("skuRole") == "specImage"), None)

    # 产品级字段（非 skuSpec）：取第一个数据行的值
    product_cols = [c for c in cols if c.get("kind") != "skuSpec"]
    data_start = 5
    product = {}
    product_keys = {}
    for c in product_cols:
        v = ws.cell(data_start, c["col"]).value
        if v is None or not str(v).strip():
            continue
        product[c["label"]] = str(v).strip() if not isinstance(v, (int, float)) else v
        product_keys[c["label"]] = c.get("key")

    # SKU 行
    skus = []
    dim_names = [c["label"] for c in dim_cols]
    dim_values = {n: [] for n in dim_names}
    for r in range(data_start, ws.max_row + 1):
        vals = {}
        for c in dim_cols:
            v = ws.cell(r, c["col"]).value
            vals[c["label"]] = str(v).strip() if v is not None and str(v).strip() else None
        if not any(vals.values()):
            continue
        row = {}
        for n in dim_names:
            row[n] = vals.get(n) or ""
            if row[n] and row[n] not in dim_values[n]:
                dim_values[n].append(row[n])
        def num(col):
            if not col:
                return None
            v = ws.cell(r, col["col"]).value
            if v is None or str(v).strip() == "":
                return None
            try:
                return float(v)
            except Exception:
                return None
        row["price"] = num(col_price)
        row["stock"] = num(col_stock)
        row["code"] = (str(ws.cell(r, col_code["col"]).value).strip()
                       if col_code and ws.cell(r, col_code["col"]).value not in (None, "") else None)
        row["barcode"] = (str(ws.cell(r, col_barcode["col"]).value).strip()
                          if col_barcode and ws.cell(r, col_barcode["col"]).value not in (None, "") else None)
        row["specImage"] = (str(ws.cell(r, col_spec_img["col"]).value).strip()
                            if col_spec_img and ws.cell(r, col_spec_img["col"]).value not in (None, "") else None)
        row["_row"] = r
        skus.append(row)

    # 预检
    warnings = []
    blocking = []
    if not skus:
        blocking.append("no_sku_rows")
    for i, s in enumerate(skus):
        for f in ("price", "stock"):
            if s.get(f) is None:
                blocking.append("sku_%d_empty_%s" % (i, f))
    codes = [s["code"] for s in skus if s.get("code")]
    if codes and len(set(codes)) != len(codes):
        warnings.append("duplicate_codes")

    allowed = load_options(wb)
    attr_check = []
    for name, vals in dim_values.items():
        attr_check.append({"dim": name, "n": len(vals)})
    # 产品级属性 vs 选项表校验
    attr_warnings = []
    for label, val in product.items():
        opts = allowed.get(label)
        if not opts:
            continue
        raw = str(val).strip()
        # 选项本身就含 /（如「成品地垫（元/块）」）→ 整串命中允许值就别再拆，
        # 否则会被拆成「成品地垫（元」「块）」并误报「不在允许值里」。
        vs = [raw] if raw in opts else [x.strip() for x in re.split(r"[,，/、]", raw) if x.strip()]
        for v in vs:
            if v not in opts:
                attr_warnings.append({"attr": label, "value": v, "hint": "不在选项_dy 允许值里，网页上可能选不中"})

    out = {
        "ok": len(blocking) == 0,
        "platform": "douyin",
        "sheet": sheet_name,
        "category": category,
        "cate_id": cate_id,
        "title": product.get("商品标题"),
        "product": product,
        "product_keys": product_keys,
        # 产品级类目属性的**完整名单 + 必填标记**（key = category_properties.<id>）。
        # build_product 用它算「模板里有、但我们没填」的属性 —— 跑完要列给用户问（门禁 47）。
        "product_fields": [{"label": c["label"], "key": c["key"], "required": c["required"]}
                           for c in cols if str(c.get("key") or "").startswith("category_properties.")],
        "dimensions": dim_names,
        "dim_values": dim_values,
        "sku_count": len(skus),
        "skus": skus,
        "sku_cols": [{"label": c["label"], "key": c["key"], "required": c["required"],
                      "skuRole": c["skuRole"], "assetPath": c["assetPath"]} for c in sku_cols],
        "asset_dirs": sorted({a for c in cols for a in (c.get("assetPath") or [])}),
        "allowed_options": allowed,
        "warnings": warnings + [w for w in attr_warnings],
        "attr_warnings": attr_warnings,
        "blocking_warnings": blocking,
    }
    print(json.dumps(out, ensure_ascii=False, indent=1))
    sys.exit(0 if out["ok"] else 2)


if __name__ == "__main__":
    main()
